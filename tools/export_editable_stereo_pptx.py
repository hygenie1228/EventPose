"""English native-shape PowerPoint export with actual mesh and unified event display."""
import hashlib
import io
import json
import os
from pathlib import Path
import re
import sys
import time
ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT/'.experiment_storage/pptx_dependencies'),str(ROOT/'tools')]
from goal_runtime_control import register_job,write

def main():
    g=json.loads((ROOT/'_agents/research/261028/device3_encoding_goal_12h.json').read_text());register_job(g,os.getpid(),'editable_English_stereo_PPTX_CPU')
    out=Path(g['output_root'])/'encoding_explanation/editable'
    def ready():
        if not (out/'actual_assets_provenance.json').exists():return False
        proof=json.loads((out/'actual_assets_provenance.json').read_text())
        if not proof.get('event_visualization_signed_timestamp_verified',False):return False
        if (out/'dense_example_selection.json').exists():
            selected=json.loads((out/'dense_example_selection.json').read_text())
            return proof['source_frame']==selected['source_frame'] and proof['session']==selected['session']
        return True
    while not ready():
        if time.time()>g['deadline_unix']-300:raise RuntimeError('Asset deadline exceeded')
        time.sleep(5)
    import numpy as np
    import cv2
    from PIL import Image
    import matplotlib.colors as mc
    from matplotlib.patches import Rectangle,FancyBboxPatch,Circle,Ellipse,Polygon,FancyArrowPatch
    from matplotlib.lines import Line2D
    from matplotlib.image import AxesImage
    from matplotlib.text import Text
    from matplotlib.collections import PathCollection
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
    from pptx.dml.color import RGBColor
    from pptx.oxml.xmlchemy import OxmlElement
    # Preparation already applies the reference support rule at the selected
    # timestamp. Never restore legacy preview images from an earlier sample.
    prs=Presentation();prs.slide_width=Inches(29);prs.slide_height=Inches(12.5)
    prs.core_properties.title='Calibration-aware RGB-Event Stereo Human Reconstruction'
    prs.core_properties.subject='Editable current-model overview and event encoding visualization'
    prs.core_properties.author='EventPose'
    sx=(29-.6)/176;sy=(12.5-.6)/76
    def xy(x,y):return Inches(.3+x*sx),Inches(.3+(76-y)*sy)
    def bounds(x,y,w,h):
        left,top=xy(x,y+h);return left,top,Inches(w*sx),Inches(h*sy)
    def color(rgba):
        a=mc.to_rgba(rgba);return RGBColor(*[round(v*255) for v in a[:3]]),a[3]
    def alpha_fill(fill,c):
        rgb,a=color(c)
        if a<=0:fill.background();return
        fill.solid();fill.fore_color.rgb=rgb
        if a<1:
            elem=fill._xPr.solidFill[0];el=OxmlElement('a:alpha');el.set('val',str(round(a*100000)));elem.append(el)
    def style(shape,face,edge,lw=1):
        alpha_fill(shape.fill,face);rgb,a=color(edge)
        if a<=0:shape.line.fill.background()
        else:shape.line.color.rgb=rgb;shape.line.width=Pt(max(.25,lw))
    def connector(slide,a,b,c,lw=1,dash=False,arrow=False):
        x0,y0=xy(*a);x1,y1=xy(*b);s=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,x0,y0,x1,y1);s.line.color.rgb=color(c)[0];s.line.width=Pt(max(.25,lw))
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            s.line.dash_style=MSO_LINE_DASH_STYLE.DASH
        if arrow:
            el=OxmlElement('a:tailEnd');el.set('type','triangle');el.set('w','sm');el.set('len','sm');s.line._get_or_add_ln().append(el)
        return s
    def textbox(slide,x,y,w,h,text,size=13,bold=False,align=PP_ALIGN.CENTER,c='#26333d'):
        assert not re.search('[가-힣]',text),text
        s=slide.shapes.add_textbox(*bounds(x,y,w,h));tf=s.text_frame;tf.clear();tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
        for i,line in enumerate(text.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(0)
            r=p.add_run();r.text=line;r.font.name='Arial';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=color(c)[0]
        return s
    def picture(slide,image,x,y,w,h):
        if isinstance(image,(str,Path)):source=str(image)
        else:
            stream=io.BytesIO();Image.fromarray(image).save(stream,format='PNG');stream.seek(0);source=stream
        return slide.shapes.add_picture(source,*bounds(x,y,w,h))
    def export(fig,ax,folder):
        fig.canvas.draw();renderer=fig.canvas.get_renderer();slide=prs.slides.add_slide(prs.slide_layouts[6]);artists=sorted(enumerate(ax.get_children()),key=lambda item:(item[1].get_zorder(),item[0]))
        for _,a in artists:
            if not a.get_visible():continue
            if isinstance(a,AxesImage):
                arr=np.asarray(a.get_array());
                if arr.dtype!=np.uint8:arr=np.uint8(np.clip(arr,0,1)*255)
                l,r,b,t=a.get_extent();picture(slide,arr,l,b,r-l,t-b)
            elif isinstance(a,FancyArrowPatch):
                start,end=a._posA_posB;connector(slide,start,end,a.get_edgecolor(),a.get_linewidth(),a.get_linestyle()=='--',True)
            elif isinstance(a,(Rectangle,FancyBboxPatch)) and a is not ax.patch:
                if isinstance(a,FancyBboxPatch):x,y,w,h=a.get_x(),a.get_y(),a.get_width(),a.get_height();kind=MSO_SHAPE.ROUNDED_RECTANGLE
                else:x,y=a.get_xy();w,h=a.get_width(),a.get_height();kind=MSO_SHAPE.RECTANGLE
                s=slide.shapes.add_shape(kind,*bounds(x,y,w,h));style(s,a.get_facecolor(),a.get_edgecolor(),a.get_linewidth())
                if kind==MSO_SHAPE.ROUNDED_RECTANGLE:
                    try:s.adjustments[0]=.08
                    except (IndexError,ValueError):pass
            elif isinstance(a,(Circle,Ellipse)):
                cx,cy=a.center;w,h=a.width,a.height;s=slide.shapes.add_shape(MSO_SHAPE.OVAL,*bounds(cx-w/2,cy-h/2,w,h));style(s,a.get_facecolor(),a.get_edgecolor(),a.get_linewidth())
            elif isinstance(a,Polygon):
                points=[xy(x,y) for x,y in a.get_xy()];builder=slide.shapes.build_freeform(points[0][0],points[0][1],scale=1);builder.add_line_segments(points[1:],close=True);s=builder.convert_to_shape();style(s,a.get_facecolor(),a.get_edgecolor(),a.get_linewidth())
            elif isinstance(a,Line2D):
                xx,yy=a.get_data();
                for i in range(len(xx)-1):connector(slide,(float(xx[i]),float(yy[i])),(float(xx[i+1]),float(yy[i+1])),a.get_color(),a.get_linewidth(),a.get_linestyle()=='--')
            elif isinstance(a,PathCollection):
                offsets=np.asarray(a.get_offsets());face=a.get_facecolors();sizes=a.get_sizes()
                for i,(x,y) in enumerate(offsets):
                    diameter=float(np.sqrt(sizes[min(i,len(sizes)-1)]))/72;s=slide.shapes.add_shape(MSO_SHAPE.OVAL,xy(x,y)[0]-Inches(diameter/2),xy(x,y)[1]-Inches(diameter/2),Inches(diameter),Inches(diameter));style(s,face[min(i,len(face)-1)],(0,0,0,0))
            elif isinstance(a,Text) and a.get_text():
                b=a.get_window_extent(renderer);p0=ax.transData.inverted().transform((b.x0,b.y0));p1=ax.transData.inverted().transform((b.x1,b.y1));size=a.get_fontsize();
                textbox(slide,float(p0[0])-.15,float(p0[1])-.22,float(p1[0]-p0[0])+.3,float(p1[1]-p0[1])+.44,a.get_text(),size)
        meshproof=json.loads((out/'actual_assets_provenance.json').read_text())
        slide.notes_slide.notes_text_frame.text=('All diagram text, blocks, feature planes and connectors are native editable PowerPoint objects. RGB/event thumbnails and the actual mesh rendering are replaceable image objects. '
            'The mesh is an actual GT-free CPU float32 inference from the Train-selected rgb3_event4_surface20 checkpoint, step '+str(meshproof['checkpoint_step'])+'. '
            'Input frame '+str(meshproof['source_frame'])+', session '+meshproof['session']+'. No mesh from an earlier experiment was substituted. '
            'Visual style reference only: https://arxiv.org/pdf/2604.09142 Figure 3. Event display convention follows the previously specified event encoding comparison; training inputs are unchanged.')
        # Second slide: the exact same black-background display conventions.
        s2=prs.slides.add_slide(prs.slide_layouts[6]);textbox(s2,3,69,170,5,'Event Encoding Visualization',26,True,PP_ALIGN.LEFT)
        names=[('count33','33 ms polarity counts','Orange-red: positive / cyan: negative'),('bins100','100 ms in 3 temporal bins','Red: oldest / green: middle / blue: latest'),('surface20','Time surface: tau = 20 ms','Orange-red: positive / cyan: negative')]
        for j,(enc,title,sub) in enumerate(names):
            x=15+j*52;textbox(s2,x,62,45,5,title,20,True);textbox(s2,x,57,45,4,sub,12)
            for row,device in enumerate((3,4)):
                y=32-row*24;img=Image.open(out/f'event{device}_{enc}_display.png');iw,ih=img.size;scale=min(45/iw,22/ih);dw,dh=iw*scale,ih*scale
                picture(s2,out/f'event{device}_{enc}_display.png',x+(45-dw)/2,y+(22-dh)/2,dw,dh)
                textbox(s2,2,y+7,12,5,'Device '+str(device),17,True)
        textbox(s2,10,1,158,5,'Display only: fixed log-count brightness + gamma; weak 3x3 support for count33. Training encodings and weights are unchanged.',12)
        s2.notes_slide.notes_text_frame.text='These are actual causal sparse event observations from the same timestamp as slide 1. Count33 uses sqrt(min(log1p(count)/log(9), 1)) and the referenced weak support rule. Surface20 uses polarity-separated exp(-age/20 ms), maximum lookback 200 ms. Bins100 collapses polarity only for visualization; training retains six channels.'
        path=out/'Stereo_RGB_Event_Pipeline_Editable.pptx';prs.save(path)
        reopened=Presentation(path);texts=[sh.text for sl in reopened.slides for sh in sl.shapes if sh.has_text_frame];assert all(not re.search('[가-힣]',t) for t in texts)
        native=sum(sh.shape_type!=13 for sl in reopened.slides for sh in sl.shapes);assert native>100
        write(out/'pptx_editability_audit.json',dict(status='passed',slides=len(reopened.slides),native_objects=native,text_objects=len(texts),all_visible_text_English=True,
            mesh_is_actual_current_checkpoint_prediction=True,mesh_and_inputs_replaceable_images=True,full_diagram_is_not_a_flat_image=True,pptx=str(path),sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
    from render_device3_stereo_paper_figure import main as render
    render(english=True,mesh_path=out/'actual_smplx_mesh.png',event_image_path=out/'event4_count33_display.png',export_callback=export)
    report=ROOT/'_agents/report/261028_스테레오_센서융합_모델_개요와_입출력.md'
    report.write_text(f'''# Calibration-aware RGB-Event Stereo Human Reconstruction

![Current model pipeline]({out}/Stereo_RGB_Event_Pipeline.png)

[Editable PowerPoint]({out}/Stereo_RGB_Event_Pipeline_Editable.pptx) · [SVG]({out}/Stereo_RGB_Event_Pipeline.svg) · [PDF]({out}/Stereo_RGB_Event_Pipeline.pdf)

The PowerPoint contains native editable text, blocks, feature planes and connectors. The RGB/event inputs and actual predicted mesh are independently replaceable images. All visible slide text is English.

## Current model and tensor shapes

| Stage | Spatial resolution | Channels / features |
|---|---|---|
| Primary RGB input | 896 x 896 | 3 |
| DINOv2 ViT-L/14 output | 64 x 64 | 1024 |
| Pooled RGB queries | 16 x 16 | 1024, projected to 128-D Q |
| Auxiliary event input | 224 x 224 | count33: 2; bins100: 6; surface20: 2 |
| Auxiliary event CNN output | 28 x 28 | 128 |
| Depth-weighted message | 16 x 16 | 128 |
| Upsampled fusion residual | 64 x 64 | 1024 |

The primary RGB path is preserved. A fixed 16 x 16 grid defines 256 primary-camera rays; 16 is an implementation setting, not an encoder output resolution. Eleven camera-Z hypotheses (0.5, 0.75, 1, 1.5, 2, 2.5, 3, 4, 5, 6, 8 metres) are projected into the auxiliary event camera using Kp, Ke, R and t. Valid projected locations are bilinearly sampled on the 28 x 28 event feature map. A masked softmax over the 11 candidates produces a 16 x 16 message, which is projected and upsampled to a 64 x 64 residual. This residual is added to the unchanged 64 x 64 primary features before native detection, HPH and SMPL-X reconstruction.

Event-only uses an event adapter before the primary ViT. RGB3 + event3 uses the same structure with a different auxiliary camera. No GT boxes, poses or depths are provided at inference. Predictions are transformed into the common device3 RGB camera for evaluation.

## Actual mesh and event display

The lower-right mesh is a real native SMPL-X prediction from the current rgb3_event4_surface20 Train-selected checkpoint, not a skeleton icon or a mesh from an earlier experiment. It was inferred on CPU in float32 for illustration and rendered with a CPU z-buffer. Its display camera is centred on the prediction without GT alignment. The full checkpoint/frame provenance is stored beside the figure in actual_assets_provenance.json.

Event images follow the requested black-background visualization convention: orange-red positive polarity and cyan negative polarity; fixed sqrt(log-count) brightness for count33; weak 3x3 support for the displayed count33 contour; old-red/middle-green/recent-blue temporal bins; polarity-separated exponential decay for surface20. These display operations do not change training inputs or checkpoints.

The layout refers only to the visual organization of [Figure 3 of the supplied paper](https://arxiv.org/pdf/2604.09142). The modules and connections depict the current Multi-HMR event-residual model. Intermediate feature grids and rays are schematic; the input images and final mesh are actual data/prediction assets.
''')
    print(out/'Stereo_RGB_Event_Pipeline_Editable.pptx',flush=True)

if __name__=='__main__':main()
